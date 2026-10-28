"""The library (``auctions/documents/``): reading files, passages, tags, search, who sees what, the pages,
and the ``/mcp/`` tools. Every model call goes to :class:`FakeProvider`; nothing here reaches a network.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import socket
import tempfile
import zipfile
from unittest import mock

from django.contrib.auth.models import User
from django.core.files.base import ContentFile
from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import Client, RequestFactory, TestCase, override_settings
from django.urls import reverse
from PIL import Image

from auctions import llm, palette_actions, palette_routes
from auctions.documents import extract as extract_module
from auctions.documents import index, pictures
from auctions.documents.extract import ExtractionError, Reader, extract
from auctions.documents.index import chunk
from auctions.documents.models import Visibility
from auctions.documents.search import can_manage, search, visible_documents
from auctions.mcp import prompts, resources, tools
from auctions.models import (
    Club,
    ClubMember,
    CopyrightNotice,
    CopyrightStrike,
    Document,
    DocumentChunk,
    DocumentFeedback,
    DocumentImageText,
    Species,
    UserData,
)
from auctions.test_support import isolated_cache

ARTICLE = """# Spawning dwarf cichlids

## Setting up

The pair of Apistogramma agassizii went into a ten gallon tank with a clay pot and soft water.

## Raising the fry

The fry ate microworms and newly hatched brine shrimp from the fifth day.
"""


class FakeProvider(llm.LLMProvider):
    """Answers every kind of call, deterministically, and counts them."""

    name = "fake"

    def __init__(self):
        super().__init__(model="fake", api_key="x")
        self.calls = {"complete": 0, "complete_json": 0, "embed": 0}
        self.transcription = "PAGE TEXT read by the vision model"
        self.details = {"topics": ["breeding", "made_up"], "title": "Dwarf cichlids", "author": "R. T.", "year": 1979}
        self.answer = {"answer": "On microworms.", "sources": [1, 99]}

    def complete(self, system, messages, tools=None, max_tokens=2000, tool_choice=""):
        self.calls["complete"] += 1
        return llm.LLMResult(text=self.transcription)

    def complete_json(self, system, messages, max_tokens=2000):
        self.calls["complete_json"] += 1
        message = messages[0]["content"]
        if "Question:" in message:
            return llm.LLMResult(data=self.answer)
        if "\nPages:\n" in message:
            return llm.LLMResult(data=self.stitch(message))
        return llm.LLMResult(data=self.details)

    #: Stand-in for the stitching model: a paragraph is about bettas, killifish, or neither.
    STORIES = {"betta": ("Spawning Bettas", "N1"), "killi": ("Killifish Notes", "N2")}

    def stitch(self, message):
        open_articles = re.findall(r"^\[(A\d+)\] (.*?), ends so far", message, re.MULTILINE)
        labels, new = {}, {}
        for paragraph, text in re.findall(r"^\[(\d+\.\d+)\] (.*)$", message, re.MULTILINE):
            story = next((name for name in self.STORIES if name in text.lower()), None)
            if story is None:
                labels[paragraph] = "skip"
                continue
            title, new_key = self.STORIES[story]
            key = next((key for key, said in open_articles if title in said), None)
            if key is None:
                key = new_key
                new[key] = {"title": title, "author": "", "year": None}
            labels[paragraph] = key
        return {"paragraphs": labels, "new": new}

    def embed(self, texts, model, dimensions=0):
        """A bag of words hashed into eight dimensions: shared words, nearby vectors."""
        self.calls["embed"] += 1
        vectors = []
        for text in texts:
            vector = [0.0] * 8
            for word in text.lower().split():
                vector[sum(map(ord, word)) % 8] += 1.0
            vectors.append([value or 0.01 for value in vector])
        return llm.EmbeddingResult(vectors=vectors, model=model)


def picture(width=400, height=300, frames=1, fmt="PNG") -> bytes:
    out = io.BytesIO()
    images = [Image.new("RGB", (width, height), (200, 200, 200 - n)) for n in range(frames)]
    images[0].save(out, fmt, save_all=frames > 1, append_images=images[1:])
    return out.getvalue()


@isolated_cache("library")
@override_settings(LLM_EMBEDDING_MODEL="fake-embedding", DOCUMENT_MODEL="fake-vision", DOCUMENT_SERVER_OCR=True)
class LibraryTestCase(TestCase):
    """A club with a member and an outsider, a fake model, and a throwaway ``DOCUMENT_ROOT``."""

    @classmethod
    def setUpClass(cls):
        cls._documents_tmp = tempfile.TemporaryDirectory()
        cls._documents_override = override_settings(DOCUMENT_ROOT=cls._documents_tmp.name)
        cls._documents_override.enable()
        super().setUpClass()

    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        cls._documents_override.disable()
        cls._documents_tmp.cleanup()

    @classmethod
    def setUpTestData(cls):
        cls.owner = User.objects.create_user("owner", "owner@example.com", "x")
        cls.member = User.objects.create_user("member", "member@example.com", "x")
        cls.outsider = User.objects.create_user("outsider", "outsider@example.com", "x")
        cls.staff = User.objects.create_superuser("staff", "staff@example.com", "x")
        # Through the cached instance, which run_as hands the tools: a queryset update wouldn't reach it.
        for user in (cls.owner, cls.member, cls.outsider, cls.staff):
            user.userdata.library_enabled = True
            user.userdata.save(update_fields=["library_enabled"])
        cls.club = Club.objects.create(name="Fish Club", abbreviation="FC")
        ClubMember.objects.create(club=cls.club, user=cls.owner, name="Owner", permission_edit_club=True)
        ClubMember.objects.create(club=cls.club, user=cls.member, name="Member")
        cls.species = Species.objects.create(
            scientific_name="Apistogramma agassizii", genus="Apistogramma", species="agassizii"
        )

    def setUp(self):
        super().setUp()
        self.provider = FakeProvider()
        llm.set_provider_override(self.provider)
        self.addCleanup(llm.set_provider_override, None)

    def document(
        self, text=ARTICLE, name="article.md", club="club", owner=None, index_it=True, visibility=None
    ) -> Document:
        """A document; by default a club-only one in the club, or a private one with ``club=None``."""
        club = self.club if club == "club" else club
        visibility = visibility or (Visibility.CLUB if club else Visibility.PRIVATE)
        document = Document(owner=owner or self.owner, club=club, original_name=name, visibility=visibility)
        document.file.save(name, ContentFile(text.encode() if isinstance(text, str) else text), save=True)
        if index_it:
            index.index_document(document.pk)
            document.refresh_from_db()
        return document


class ChunkTests(TestCase):
    def test_passages_split_at_headings_and_point_back_into_the_text(self):
        passages = chunk(ARTICLE)
        self.assertEqual(
            [p.heading for p in passages],
            [
                "Spawning dwarf cichlids",
                "Spawning dwarf cichlids > Setting up",
                "Spawning dwarf cichlids > Raising the fry",
            ],
        )
        for passage in passages:
            self.assertEqual(ARTICLE[passage.start : passage.end], passage.text)

    def test_a_long_paragraph_is_split_without_losing_text(self):
        text = " ".join(f"word{n}." for n in range(2000))
        passages = chunk(text)
        self.assertGreater(len(passages), 3)
        self.assertTrue(all(len(p.text) <= index.TARGET_CHARS for p in passages))
        self.assertEqual("".join(text[p.start : p.end] for p in passages).replace(" ", ""), text.replace(" ", ""))

    def test_page_markers_set_the_page_and_are_not_text(self):
        passages = chunk("<!-- page 1 -->\n\nfirst page\n\n<!-- page 2 -->\n\nsecond page")
        self.assertEqual([(p.page, p.text) for p in passages], [(1, "first page"), (2, "second page")])

    def test_a_split_table_repeats_its_header(self):
        rows = "\n".join(f"| fish {n} | {n} |" for n in range(200))
        passages = chunk(f"| name | count |\n|---|---|\n{rows}")
        self.assertGreater(len(passages), 1)
        self.assertTrue(all(p.text.startswith("| name | count |") for p in passages))


class ExtractTests(LibraryTestCase):
    def test_markdown_text_and_csv_come_through_markitdown(self):
        self.assertIn("clay pot", extract(ARTICLE.encode(), "a.md").text)
        self.assertIn("guppy", extract(b"plain words about a guppy", "a.txt").text)
        self.assertIn("Betta", extract(b"name,count\nBetta,3\n", "a.csv").text)

    def test_html_is_converted_without_touching_the_network(self):
        page = b'<html><body><h1>Care sheet</h1><img src="http://169.254.169.254/latest"></body></html>'

        def refuse(*args, **kwargs):
            message = "tried to open a network connection"
            raise AssertionError(message)

        with mock.patch.object(socket.socket, "connect", refuse):
            self.assertIn("Care sheet", extract(page, "page.html").text)

    def test_a_picture_is_read_by_the_vision_model_once_per_document(self):
        document = self.document(index_it=False)
        first = extract(picture(), "scan.jpg", Reader(document))
        self.assertEqual(first.text, self.provider.transcription)
        extract(picture(), "scan.jpg", Reader(document))
        self.assertEqual(self.provider.calls["complete"], 1, "the second read should come from the cache")
        document.delete()
        self.assertFalse(DocumentImageText.objects.exists(), "transcriptions go with their document")

    def test_frames_past_the_allowance_are_not_decoded_and_are_noted(self):
        result = extract(picture(frames=3, fmt="TIFF"), "scan.tif", Reader(allowance=2))
        self.assertEqual(result.text.count("<!-- page"), 2)
        self.assertIn("1 pictures or pages weren't read", result.notes[0])

    def test_every_frame_of_a_multipage_tiff_is_a_page(self):
        text = extract(picture(frames=2, fmt="TIFF"), "scan.tif").text
        self.assertIn("<!-- page 2 -->", text)

    def test_a_tiny_picture_is_refused(self):
        with self.assertRaises(ExtractionError):
            extract(picture(20, 20), "icon.png")

    def test_a_scanned_pdf_is_read_page_by_page(self):
        out = io.BytesIO()
        Image.new("RGB", (850, 1100), "white").save(out, "PDF")
        text = extract(out.getvalue(), "scan.pdf").text
        self.assertIn("<!-- page 1 -->", text)
        self.assertIn(self.provider.transcription, text)

    def test_a_text_pdf_goes_to_markitdown(self):
        from reportlab.pdfgen import canvas

        out = io.BytesIO()
        page = canvas.Canvas(out)
        page.drawString(72, 720, "The killifish eggs were stored in damp peat for three months.")
        page.save()
        text = extract(out.getvalue(), "typed.pdf").text
        self.assertIn("damp peat", text)
        self.assertEqual(self.provider.calls["complete"], 0)

    def test_iphone_photos_open(self):
        """HEIC is what an iPhone saves a photographed page as; Pillow only knows it once the plugin loads."""
        extract_module.frames_as_jpeg(b"")
        self.assertIn(".heic", Image.registered_extensions())

    def test_a_typed_pdf_with_a_full_page_picture_is_read_as_a_scan(self):
        """A scan with an old OCR layer: plenty of text, but one picture covers the page."""
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfgen import canvas

        out = io.BytesIO()
        page = canvas.Canvas(out, pagesize=(600, 800))
        page.drawImage(ImageReader(io.BytesIO(picture(600, 800))), 0, 0, width=600, height=800)
        page.drawString(72, 720, "Bad OCR layer from 1998 " * 5)
        page.save()
        text = extract(out.getvalue(), "scan-with-ocr.pdf").text
        self.assertIn(self.provider.transcription, text)

    def test_pictures_inside_a_powerpoint_are_read(self):
        """markitdown asks its OpenAI-shaped client for a caption; ours answers. If markitdown changes
        how it asks, this is what notices.
        """
        from pptx import Presentation

        deck = Presentation()
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(picture()), 0, 0)
        out = io.BytesIO()
        deck.save(out)
        self.assertIn(self.provider.transcription, extract(out.getvalue(), "talk.pptx").text)

    def test_failed_vision_reads_are_counted_not_raised(self):
        self.provider.complete = mock.Mock(side_effect=llm.LLMError("down"))
        reader = Reader()
        self.assertEqual(reader.read(picture()), "")
        self.assertEqual(reader.failed, 1)

    def test_the_allowance_caps_vision_reads(self):
        reader = Reader(allowance=0)
        self.assertEqual(reader.read(picture()), "")
        self.assertEqual(reader.skipped, 1)

    def test_a_zip_bomb_is_refused_before_it_is_opened(self):
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as archive:
            for n in range(extract_module.MAX_ZIP_ENTRIES + 1):
                archive.writestr(f"{n}.txt", "x")
        with self.assertRaises(ExtractionError):
            extract(out.getvalue(), "bomb.zip")

    def test_markitdown_still_has_the_docx_picture_hook(self):
        """extract.py overrides it; a markitdown upgrade that renames it would silently stop reading pictures."""
        from markitdown.converters import DocxConverter

        self.assertTrue(callable(getattr(DocxConverter, "_image_to_html", None)))


class IndexTests(LibraryTestCase):
    def test_a_document_is_read_chunked_embedded_and_tagged(self):
        document = self.document()
        self.assertEqual(document.status, Document.READY, document.error)
        self.assertEqual(document.chunks.count(), 3)
        self.assertFalse(document.chunks.filter(embedding=None).exists())
        self.assertEqual(list(document.species.all()), [self.species])
        self.assertEqual(document.topics, ["breeding"], "a topic not on the list is dropped")
        self.assertEqual((document.title, document.author, document.year), ("Dwarf cichlids", "R. T.", 1979))

    def test_what_people_typed_is_never_overwritten(self):
        document = self.document(index_it=False)
        Document.objects.filter(pk=document.pk).update(title="Their title", year=1981)
        index.index_document(document.pk)
        document.refresh_from_db()
        self.assertEqual((document.title, document.year), ("Their title", 1981))

    def test_rechunking_keeps_vectors_for_unchanged_passages(self):
        document = self.document()
        embeds = self.provider.calls["embed"]
        Document.objects.filter(pk=document.pk).update(chunker_version="old")
        index.index_document(document.pk)
        self.assertEqual(self.provider.calls["embed"], embeds)
        self.assertFalse(document.chunks.filter(embedding=None).exists())

    def test_a_new_parser_rereads_the_file_unless_the_text_was_corrected(self):
        document = self.document()
        Document.objects.filter(pk=document.pk).update(
            parser_version="old", text="Corrected by hand", text_edited=True, chunker_version=""
        )
        index.index_document(document.pk)
        document.refresh_from_db()
        self.assertEqual(document.text, "Corrected by hand")
        self.assertEqual(document.chunks.get().text, "Corrected by hand")
        self.assertNotIn(document, index.stale_documents(), "a corrected text is never stale for a new parser")

    def test_read_it_again_asks_the_model_again(self):
        document = self.document(text=picture(), name="scan.png")
        self.assertEqual(self.provider.calls["complete"], 1)
        with mock.patch("auctions.tasks.index_document.delay", side_effect=index.index_document):
            with self.captureOnCommitCallbacks(execute=True):
                index.queue(document, reread=True)
        self.assertEqual(self.provider.calls["complete"], 2, "a saved reading would hand back the same bad text")

    def test_rechunking_does_not_ask_the_model_for_details_again(self):
        document = self.document()
        Document.objects.filter(pk=document.pk).update(author="", chunker_version="old")
        asked = self.provider.calls["complete_json"]
        index.index_document(document.pk)
        self.assertEqual(self.provider.calls["complete_json"], asked)

    def test_reread_discards_corrections(self):
        document = self.document()
        Document.objects.filter(pk=document.pk).update(text="Corrected by hand", text_edited=True)
        with mock.patch("auctions.tasks.index_document.delay", side_effect=index.index_document):
            with self.captureOnCommitCallbacks(execute=True):
                index.queue(document, reread=True)
        document.refresh_from_db()
        self.assertIn("clay pot", document.text)
        self.assertFalse(document.text_edited)

    def test_an_unreadable_file_fails_with_a_reason(self):
        document = self.document(text=picture(10, 10), name="dot.png")
        self.assertEqual(document.status, Document.FAILED)
        self.assertIn("too small", document.error)

    def test_stale_documents_follow_the_embedding_model(self):
        document = self.document()
        self.assertNotIn(document, index.stale_documents())
        with override_settings(LLM_EMBEDDING_MODEL="another-model"):
            self.assertIn(document, index.stale_documents())

    @override_settings(LLM_EMBEDDING_MODEL="")
    def test_no_embedding_model_means_keyword_search_only(self):
        document = self.document()
        self.assertEqual(document.status, Document.READY)
        self.assertEqual(self.provider.calls["embed"], 0)
        hits, _total = search(self.owner, "microworms")
        self.assertEqual(hits[0].document, document)

    def test_requeue_picks_up_documents_nobody_is_reading_or_waiting_for(self):
        document = self.document(index_it=False)
        with mock.patch("auctions.tasks.index_document.delay") as delay:
            self.assertEqual(index.requeue_stuck(), 1)
            self.assertEqual(index.requeue_stuck(), 0, "now marked queued, so not queued twice")
        delay.assert_called_once_with(document.pk)

    def test_a_document_that_keeps_killing_the_worker_fails(self):
        document = self.document(index_it=False)
        Document.objects.filter(pk=document.pk).update(status=Document.PROCESSING, attempts=index.MAX_ATTEMPTS)
        with mock.patch("auctions.tasks.index_document.delay") as delay:
            index.requeue_stuck()
        delay.assert_not_called()
        document.refresh_from_db()
        self.assertEqual(document.status, Document.FAILED)

    def test_a_running_document_is_not_run_twice(self):
        document = self.document(index_it=False)
        from django.core.cache import cache

        cache.add(index._lock(document.pk), 1)
        self.assertFalse(index.index_document(document.pk), "the task retries on False")

    def test_an_edit_saved_during_a_read_is_kept(self):
        document = self.document()
        real_tag = index.tag

        def someone_edits_meanwhile(doc, **kwargs):
            Document.objects.filter(pk=doc.pk).update(
                title="Typed meanwhile", text="Edited meanwhile", chunker_version=""
            )
            real_tag(doc, **kwargs)

        Document.objects.filter(pk=document.pk).update(title="", chunker_version="old")
        with mock.patch.object(index, "tag", someone_edits_meanwhile):
            index.index_document(document.pk)
        document.refresh_from_db()
        self.assertEqual((document.title, document.text), ("Typed meanwhile", "Edited meanwhile"))
        self.assertEqual(document.chunker_version, "", "left stale, so the queued run re-cuts the edited text")


class VisibilityTests(LibraryTestCase):
    def test_club_documents_are_for_members_and_personal_ones_for_their_owner(self):
        club_document = self.document(index_it=False)
        private = self.document(club=None, name="mine.md", index_it=False)
        self.assertEqual(set(visible_documents(self.owner)), {club_document, private})
        self.assertEqual(set(visible_documents(self.member)), {club_document})
        self.assertEqual(set(visible_documents(self.outsider)), set())
        self.assertEqual(set(visible_documents(self.staff)), set(), "being a superuser is not membership")

    def test_removed_documents_are_hidden_from_everyone(self):
        document = self.document(index_it=False)
        Document.objects.filter(pk=document.pk).update(removed=True)
        self.assertFalse(visible_documents(self.owner).exists())

    def test_managing_is_the_uploader_or_the_clubs_editors(self):
        document = self.document(index_it=False, owner=self.member)
        self.assertTrue(can_manage(self.member, document))
        self.assertTrue(can_manage(self.owner, document), "owner has permission_edit_club")
        self.assertFalse(can_manage(self.outsider, document))

    def test_search_never_returns_what_the_caller_cannot_see(self):
        self.document()
        hits, _total = search(self.outsider, "microworms brine shrimp")
        self.assertEqual(hits, [])


class SearchTests(LibraryTestCase):
    def test_keyword_and_vector_rankings_are_fused(self):
        document = self.document()
        self.document(text="# Plants\n\nJava fern grows on driftwood.", name="plants.md")
        hits, total = search(self.member, "what did the fry eat? microworms")
        self.assertEqual(hits[0].document, document)
        self.assertIn("microworms", hits[0].chunk.text)
        self.assertGreaterEqual(total, 1)

    def test_the_written_answer_cites_only_real_passages(self):
        from auctions.documents.search import answer

        self.document()
        hits, _total = search(self.member, "microworms")
        result = answer(self.member, "what did the fry eat", hits)
        self.assertEqual(result["answer"], "On microworms.")
        self.assertEqual([number for number, _hit in result["sources"]], [1])

    def test_passages_reach_the_model_fenced(self):
        from auctions.documents.search import answer

        self.document()
        hits, _total = search(self.member, "microworms")
        with mock.patch.object(self.provider, "complete_json", wraps=self.provider.complete_json) as sent:
            answer(self.member, "what did the fry eat", hits)
        self.assertIn(palette_actions.UNTRUSTED_OPEN, sent.call_args.args[1][0]["content"])


class PageTests(LibraryTestCase):
    def test_upload_saves_and_queues_the_document(self):
        self.client.force_login(self.owner)
        upload = SimpleUploadedFile("care.txt", b"Feed daphnia twice a day.")
        with self.captureOnCommitCallbacks(execute=False) as callbacks:
            response = self.client.post(reverse("library_add"), {"file": upload, "club": self.club.pk})
        document = Document.objects.get()
        self.assertRedirects(response, document.get_absolute_url())
        self.assertEqual((document.original_name, document.club, document.owner), ("care.txt", self.club, self.owner))
        self.assertEqual(len(callbacks), 1)
        self.assertNotIn("care", document.file.name, "the stored name is random")

    def test_several_files_make_several_documents_and_a_duplicate_is_skipped(self):
        self.client.force_login(self.owner)
        self.client.post(reverse("library_add"), {"file": SimpleUploadedFile("old.txt", b"already here")})
        files = [
            SimpleUploadedFile("one.txt", b"first"),
            SimpleUploadedFile("two.txt", b"second"),
            SimpleUploadedFile("again.txt", b"already here"),
        ]
        response = self.client.post(
            reverse("library_add"), {"file": files, "title": "Ignored", "author": "Ed."}, follow=True
        )
        self.assertEqual(Document.objects.count(), 3)
        self.assertEqual(set(Document.objects.filter(author="Ed.").values_list("title", flat=True)), {""})
        self.assertContains(response, "again.txt")

    def test_arriving_from_a_club_preselects_it(self):
        self.client.force_login(self.owner)
        response = self.client.get(reverse("library_add"), {"club": self.club.slug})
        self.assertEqual(response.context["form"].initial, {"visibility": Visibility.PUBLIC, "club": self.club.pk})

    def test_the_same_file_twice_is_refused(self):
        self.client.force_login(self.owner)
        for _ in range(2):
            response = self.client.post(reverse("library_add"), {"file": SimpleUploadedFile("a.txt", b"same bytes")})
        self.assertEqual(Document.objects.count(), 1)
        self.assertContains(response, "already in the library")

    def test_a_club_needs_permission_to_file_under(self):
        self.client.force_login(self.member)
        self.client.post(reverse("library_add"), {"file": SimpleUploadedFile("a.txt", b"x"), "club": self.club.pk})
        self.assertEqual(Document.objects.get().club, None, "no club they may file under, so it is theirs alone")

    def test_pages_and_the_file_are_for_members_only(self):
        document = self.document()
        self.client.force_login(self.member)
        self.assertContains(self.client.get(document.get_absolute_url()), "clay pot")
        response = self.client.get(reverse("document_file", args=[document.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/octet-stream")
        self.assertEqual(response["X-Content-Type-Options"], "nosniff")
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(document.get_absolute_url()).status_code, 404)
        self.assertEqual(self.client.get(reverse("document_file", args=[document.pk])).status_code, 404)

    def test_search_page_and_answer(self):
        self.document()
        self.client.force_login(self.member)
        self.assertContains(self.client.get(reverse("library_search"), {"q": "microworms"}), "Spawning dwarf cichlids")
        self.assertContains(self.client.get(reverse("library_answer"), {"q": "microworms"}), "On microworms.")

    def test_only_managers_edit_or_delete(self):
        document = self.document()
        self.client.force_login(self.member)
        self.assertEqual(self.client.post(reverse("document_delete", args=[document.pk])).status_code, 404)
        self.client.force_login(self.owner)
        self.client.post(reverse("document_delete", args=[document.pk]))
        self.assertFalse(Document.objects.exists())

    def test_saving_without_touching_the_text_leaves_it_alone(self):
        document = self.document()
        self.client.force_login(self.owner)
        crlf = document.text.replace("\n", "\r\n")
        data = {"title": "T", "author": "", "year": "1977", "club": self.club.pk, "topics": ["plants"], "text": crlf}
        data["visibility"] = "club"
        with self.captureOnCommitCallbacks(execute=False) as callbacks:
            self.client.post(reverse("document_edit", args=[document.pk]), data)
        document.refresh_from_db()
        self.assertEqual(document.year, 1977)
        self.assertFalse(document.text_edited)
        self.assertEqual(callbacks, [], "nothing to re-index")

    def test_a_superuser_member_can_open_the_edit_page(self):
        ClubMember.objects.create(club=self.club, user=self.staff, name="Staff")
        document = self.document(index_it=False)
        self.client.force_login(self.staff)
        self.assertEqual(self.client.get(reverse("document_edit", args=[document.pk])).status_code, 200)

    def test_a_taken_down_file_cannot_come_back(self):
        document = self.document(index_it=False)
        Document.objects.filter(pk=document.pk).update(removed=True, sha256=hashlib.sha256(b"x").hexdigest())
        self.client.force_login(self.member)
        response = self.client.post(reverse("library_add"), {"file": SimpleUploadedFile("again.txt", b"x")})
        self.assertContains(response, "taken down")
        self.assertEqual(Document.objects.count(), 1)

    def test_correcting_the_text_marks_it_edited_and_reindexes(self):
        document = self.document()
        self.client.force_login(self.owner)
        data = {"title": "T", "author": "", "year": "", "club": self.club.pk, "topics": ["plants"], "text": "Fixed"}
        data["visibility"] = "club"
        # The patch outlives the on-commit callbacks, or they would reach the real broker.
        with mock.patch("auctions.tasks.index_document.delay", side_effect=index.index_document):
            with self.captureOnCommitCallbacks(execute=True):
                self.client.post(reverse("document_edit", args=[document.pk]), data)
        document.refresh_from_db()
        self.assertTrue(document.text_edited)
        self.assertEqual(document.chunks.get().text, "Fixed")

    def test_a_reader_can_report_a_problem(self):
        document = self.document()
        self.client.force_login(self.member)
        self.client.post(reverse("document_feedback", args=[document.pk]), {"reason": "unreadable", "note": "p2"})
        self.assertEqual(DocumentFeedback.objects.get().document, document)
        self.client.force_login(self.owner)
        self.assertContains(self.client.get(document.get_absolute_url()), "The text is garbled or missing")


class McpToolTests(LibraryTestCase):
    def run_as(self, user, name, **params):
        request = RequestFactory().post("/mcp/")
        request.user = user
        request.palette_page = {}
        return palette_actions.run_action(request, name, params)

    def test_search_and_read(self):
        document = self.document()
        found = self.run_as(self.member, "search_documents", query="microworms")
        self.assertTrue(found["found"])
        self.assertEqual(found["passages"][0]["document"], document.pk)
        self.assertIn(palette_actions.UNTRUSTED_OPEN, found["passages"][0]["text"])
        read = self.run_as(self.member, "read_document", document=str(document.pk), start=0, length=20)
        self.assertEqual((read["start"], read["end"], read["length"]), (0, 20, len(document.text)))
        self.assertIn("start=20", read["summary"])

    def test_filters(self):
        self.document()
        self.assertTrue(self.run_as(self.member, "search_documents", query="pot", topic="breeding")["found"])
        self.assertFalse(self.run_as(self.member, "search_documents", query="pot", topic="ponds")["found"])
        self.assertIn("error", self.run_as(self.member, "search_documents", query="pot", topic="astrology"))
        self.assertTrue(
            self.run_as(self.member, "search_documents", query="pot", species="Apistogramma agassizii")["found"]
        )

    def test_outsiders_get_nothing(self):
        document = self.document()
        self.assertFalse(self.run_as(self.outsider, "search_documents", query="microworms")["found"])
        self.assertIn("error", self.run_as(self.outsider, "read_document", document=str(document.pk)))

    def test_an_agent_adds_its_own_transcription(self):
        text = "# Killifish eggs\n\n<!-- page 1 -->\n\nStored in peat. Species: Nothobranchius rachovii"
        result = self.run_as(
            self.owner, "add_document", title="Killifish eggs", text=text, club=self.club.slug, topics=["Breeding"]
        )
        self.assertTrue(result.get("ok"), result)
        document = Document.objects.get()
        self.assertEqual((document.club, document.text, document.topics), (self.club, text, ["breeding"]))
        self.assertTrue(document.original_name.endswith(".md"))
        with document.file.open("rb") as handle:
            self.assertEqual(handle.read().decode(), text, "the stored file is the transcription")
        again = self.run_as(self.owner, "add_document", title="Killifish eggs", text=text, club=self.club.slug)
        self.assertIn("already in the library", again["error"])

    def test_adding_to_a_club_needs_permission_and_says_so(self):
        result = self.run_as(self.member, "add_document", title="T", text="words", club=self.club.slug)
        self.assertIn("can't add documents", result["error"])
        self.assertFalse(Document.objects.exists(), "never quietly filed privately instead")
        self.assertIn("error", self.run_as(self.owner, "add_document", title="T", text="w", topics="astrology"))

    def test_list_documents_counts_reports_only_for_keepers(self):
        document = self.document()
        DocumentFeedback.objects.create(document=document, user=self.member, reason="unreadable")
        mine = self.run_as(self.owner, "list_documents", club=self.club.slug)
        self.assertEqual(mine["documents"][0]["open_reports"], 1)
        theirs = self.run_as(self.member, "list_documents")
        self.assertNotIn("open_reports", theirs["documents"][0])
        self.assertEqual(self.run_as(self.outsider, "list_documents")["documents"], [])

    def test_update_and_delete_go_through_the_same_permission(self):
        document = self.document()
        self.assertIn("error", self.run_as(self.member, "update_document", document=str(document.pk), year=1975))
        result = self.run_as(self.owner, "update_document", document=str(document.pk), year=1975, topics=["Ponds"])
        self.assertTrue(result.get("ok"), result)
        document.refresh_from_db()
        self.assertEqual((document.year, document.topics), (1975, ["ponds"]))
        self.assertFalse(document.text_edited)
        self.assertIn("error", self.run_as(self.member, "delete_document", document=str(document.pk)))
        self.assertTrue(self.run_as(self.owner, "delete_document", document=str(document.pk)).get("ok"))
        self.assertFalse(Document.objects.exists())


class ResourceTests(LibraryTestCase):
    """``document://{n}``: a citation a client can open, never a list of anybody's papers."""

    def request_as(self, user):
        request = RequestFactory().post("/mcp/")
        request.user = user
        return request

    def test_a_search_cites_each_document_once_as_a_resource_link(self):
        document = self.document()
        result = tools.call_tool(self.request_as(self.member), "search_documents", {"query": "microworms"})
        links = [block["uri"] for block in result["content"] if block["type"] == "resource_link"]
        self.assertEqual(links.count(f"document://{document.pk}"), 1)

    def test_reading_one_never_links_to_itself(self):
        document = self.document()
        result = tools.call_tool(self.request_as(self.member), "read_document", {"document": str(document.pk)})
        links = [block["uri"] for block in result["content"] if block["type"] == "resource_link"]
        self.assertNotIn(f"document://{document.pk}", links)

    def test_it_reads_as_the_tool_does_with_the_same_permission(self):
        document = self.document()
        uri = f"document://{document.pk}"
        read = resources.read(self.request_as(self.member), uri)
        self.assertEqual(json.loads(read["text"])["document"], document.pk)
        refused = resources.read(self.request_as(self.outsider), uri)
        self.assertIn("error", json.loads(refused["text"]))

    def test_it_asks_for_as_much_as_read_document_gives(self):
        template, arguments = resources.match("document://12")
        self.assertEqual((template.action, arguments), ("read_document", {"document": "12"}))
        self.assertEqual(template.extra["length"], palette_actions.READ_DOCUMENT_MAX_CHARS)

    def test_only_somebody_with_the_library_is_shown_the_template(self):
        def offered(user):
            return {row["uriTemplate"] for row in resources.template_descriptors(user)}

        self.assertIn("document://{document}", offered(self.owner))
        self.owner.userdata.library_enabled = False
        self.owner.userdata.save(update_fields=["library_enabled"])
        self.assertNotIn("document://{document}", offered(self.owner))
        self.assertNotIn("document://{document}", offered(None))


class LifecycleTests(LibraryTestCase):
    def test_deleting_a_document_deletes_its_file(self):
        document = self.document(index_it=False)
        storage, name = document.file.storage, document.file.name
        with self.captureOnCommitCallbacks(execute=True):
            document.delete()
        self.assertFalse(storage.exists(name))

    def test_account_deletion_keeps_the_clubs_documents(self):
        from auctions.account_deletion import _delete_personal_rows

        club_document = self.document(index_it=False)
        self.document(club=None, name="mine.md", index_it=False)
        _delete_personal_rows(self.owner)
        self.assertEqual(list(Document.objects.all()), [club_document])
        club_document.refresh_from_db()
        self.assertIsNone(club_document.owner)

    def test_a_copyright_notice_takes_the_document_down(self):
        from auctions import dmca
        from auctions.views.moderation import _document_from_urls

        document = self.document(index_it=False)
        self.assertEqual(_document_from_urls(f"https://example.com/library/{document.pk}/"), document)
        notice = CopyrightNotice.objects.create(
            name="R", email="r@example.com", address="a", work="w", material="m", document=document
        )
        dmca.take_down(notice)
        document.refresh_from_db()
        self.assertTrue(document.removed)
        self.assertTrue(document.file.storage.exists(document.file.name), "kept, for a counter-notice")
        self.assertEqual(CopyrightStrike.objects.filter(user=self.owner).count(), 1)


class ChunkRowTests(LibraryTestCase):
    def test_chunk_rows_store_offsets_into_the_text(self):
        document = self.document()
        for row in DocumentChunk.objects.filter(document=document):
            self.assertEqual(document.text[row.start : row.end], row.text)


class PublicByDefaultTests(LibraryTestCase):
    def test_an_upload_is_public_unless_said_otherwise(self):
        self.client.force_login(self.owner)
        self.client.post(reverse("library_add"), {"file": SimpleUploadedFile("care.txt", b"Feed daphnia.")})
        document = Document.objects.get()
        self.assertEqual(document.visibility, Visibility.PUBLIC)
        self.assertIn(document, visible_documents(self.outsider))

    def test_club_only_needs_a_club_and_hides_it_from_outsiders(self):
        self.client.force_login(self.owner)
        response = self.client.post(
            reverse("library_add"), {"file": SimpleUploadedFile("a.txt", b"a"), "visibility": Visibility.CLUB}
        )
        self.assertContains(response, "Choose the club")
        self.client.post(
            reverse("library_add"),
            {"file": SimpleUploadedFile("a.txt", b"a"), "visibility": Visibility.CLUB, "club": self.club.pk},
        )
        document = Document.objects.get()
        self.assertIn(document, visible_documents(self.member))
        self.assertNotIn(document, visible_documents(self.outsider))

    def test_private_is_only_the_owner_even_in_a_club(self):
        document = self.document(index_it=False, visibility=Visibility.PRIVATE)
        self.assertEqual(list(visible_documents(self.member)), [])
        self.assertEqual(list(visible_documents(self.owner)), [document])


def _page_texts():
    """Three pages: a betta article running on to page 2, a killifish one starting there, and an advert."""
    return [
        "SPAWNING BETTAS by J. Smith\n\nThe betta male built a bubble nest under a leaf.",
        "The betta fry were fed infusoria for a week.\n\nKILLIFISH NOTES\n\nKillifish eggs went into damp peat.",
        "The killifish eggs hatched after three months in peat.\n\nJoe's Fish Shop, open Saturdays, call 555-0100",
    ]


class BatchTests(LibraryTestCase):
    def batch(self, texts=None, visibility=Visibility.PUBLIC):
        from auctions.models import BatchPage, DocumentBatch

        batch = DocumentBatch.objects.create(owner=self.owner, club=self.club, visibility=visibility)
        for number, text in enumerate(texts or _page_texts(), 1):
            page = BatchPage(batch=batch, original_name=f"scan{number}.jpg", read=True, text=text)
            page.file.save(f"scan{number}.jpg", ContentFile(picture()), save=True)
        return batch

    def run_batch(self, batch):
        from auctions.documents import stitch

        with mock.patch("auctions.tasks.index_document.delay", side_effect=index.index_document):
            with self.captureOnCommitCallbacks(execute=True):
                while stitch.process(batch.pk) == stitch.MORE:
                    pass
        batch.refresh_from_db()

    def test_pages_become_one_document_per_article_across_page_breaks(self):
        batch = self.batch()
        self.run_batch(batch)
        self.assertEqual(batch.status, batch.DONE, batch.error)
        documents = {document.title: document for document in batch.documents.all()}
        self.assertEqual(set(documents), {"Spawning Bettas", "Killifish Notes"})
        bettas, killifish = documents["Spawning Bettas"], documents["Killifish Notes"]
        self.assertEqual([page.position for page in bettas.pages.order_by("position")], [1, 2])
        self.assertEqual([page.position for page in killifish.pages.order_by("position")], [2, 3])
        self.assertIn("<!-- page 2 -->", bettas.text)
        self.assertIn("infusoria", bettas.text)
        self.assertNotIn("Killifish", bettas.text)
        self.assertNotIn("555-0100", killifish.text, "the advert was skipped")
        self.assertEqual(
            (bettas.visibility, bettas.club, bettas.status), (Visibility.PUBLIC, self.club, Document.READY)
        )
        self.assertEqual(bettas.topics, ["breeding"], "filed under topics, keeping the stitcher's title")
        self.assertEqual(bettas.described_with, "fake-vision")

    def test_an_article_stays_open_across_windows(self):
        from auctions.documents import stitch

        texts = [f"The betta part {number} of a long article." for number in range(1, 8)]
        batch = self.batch(texts)
        with mock.patch.object(stitch, "WINDOW", 2), mock.patch.object(stitch, "DECIDED", 1):
            self.run_batch(batch)
        self.assertEqual(batch.documents.count(), 1)
        self.assertEqual(batch.documents.get().pages.count(), 7)

    def test_pages_are_put_in_file_name_order(self):
        from auctions.documents import stitch

        batch = self.batch(["a", "b", "c"])
        batch.pages.filter(original_name="scan1.jpg").update(original_name="scan10.jpg")
        stitch.put_in_order(batch)
        names = list(batch.pages.order_by("position").values_list("original_name", flat=True))
        self.assertEqual(names, ["scan2.jpg", "scan3.jpg", "scan10.jpg"])

    def test_a_model_failure_stops_where_it_was_and_carries_on(self):
        from auctions.documents import stitch

        batch = self.batch([f"The betta part {number}." for number in range(1, 8)])
        real = self.provider.complete_json
        calls = {"n": 0}

        def fails_second_time(*args, **kwargs):
            calls["n"] += 1
            if calls["n"] == 2:
                message = "down"
                raise llm.LLMError(message)
            return real(*args, **kwargs)

        with mock.patch.object(stitch, "WINDOW", 2), mock.patch.object(stitch, "DECIDED", 1):
            with mock.patch.object(self.provider, "complete_json", side_effect=fails_second_time):
                self.run_batch(batch)
            self.assertEqual(batch.status, batch.FAILED)
            self.assertEqual(batch.stitched_through, 1)
            self.client.force_login(self.owner)
            self.client.post(reverse("document_batch", args=[batch.pk]), {"action": "finish"})
            batch.refresh_from_db()
            self.assertEqual(batch.status, batch.STITCHING, "carries on from where it stopped")
            self.run_batch(batch)
        self.assertEqual(batch.documents.get().pages.count(), 7)

    def test_unread_pages_are_read_then_stitched(self):
        batch = self.batch(["x"])
        batch.pages.update(read=False, text="")
        self.provider.transcription = "The betta male built his bubble nest under a floating leaf."
        self.run_batch(batch)
        page = batch.pages.get()
        self.assertEqual((page.read, page.text, page.read_with), (True, self.provider.transcription, "fake-vision"))
        self.assertEqual(batch.documents.get().title, "Spawning Bettas")

    def test_upload_in_pages_mode_makes_a_batch_and_a_pdf_is_one_page_each(self):
        from reportlab.pdfgen import canvas

        out = io.BytesIO()
        pdf = canvas.Canvas(out)
        for _ in range(2):
            pdf.drawString(72, 720, "The betta nest " * 10)
            pdf.showPage()
        pdf.save()
        self.client.force_login(self.owner)
        files = [SimpleUploadedFile("issue.pdf", out.getvalue()), SimpleUploadedFile("loose.png", picture())]
        with self.captureOnCommitCallbacks(execute=False):
            response = self.client.post(reverse("library_add"), {"file": files, "mode": "pages", "more_coming": "on"})
        batch = self.owner.document_batches.get()
        self.assertRedirects(response, batch.get_absolute_url())
        self.assertEqual(batch.status, batch.WAITING)
        self.assertEqual(sorted(batch.pages.values_list("pdf_page", flat=True), key=str), [0, 1, None])
        pdf_pages = batch.pages.exclude(pdf_page=None)
        self.assertEqual(len({page.file.name for page in pdf_pages}), 1, "a PDF's pages share its file")

    def test_word_files_are_not_pages(self):
        self.client.force_login(self.owner)
        response = self.client.post(
            reverse("library_add"), {"file": SimpleUploadedFile("talk.docx", b"x"), "mode": "pages"}
        )
        self.assertContains(response, "Only pictures and PDFs can be pages")

    def test_scans_are_shown_to_whoever_can_read_the_article(self):
        batch = self.batch(visibility=Visibility.CLUB)
        self.run_batch(batch)
        page = batch.pages.get(position=1)
        self.client.force_login(self.member)
        self.assertEqual(self.client.get(reverse("batch_page", args=[page.pk])).status_code, 200)
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(reverse("batch_page", args=[page.pk])).status_code, 404)

    def test_put_together_again_replaces_the_articles(self):
        batch = self.batch()
        self.run_batch(batch)
        first = set(batch.documents.values_list("pk", flat=True))
        self.client.force_login(self.owner)
        with self.captureOnCommitCallbacks(execute=False):
            self.client.post(reverse("document_batch", args=[batch.pk]), {"action": "again"})
        self.run_batch(batch)
        self.assertEqual(batch.documents.count(), 2)
        self.assertFalse(first & set(batch.documents.values_list("pk", flat=True)))

    def test_a_stitched_article_is_not_read_again_from_a_file_it_does_not_have(self):
        batch = self.batch()
        self.run_batch(batch)
        document = batch.documents.first()
        Document.objects.filter(pk=document.pk).update(parser_version="old")
        index.index_document(document.pk)
        document.refresh_from_db()
        self.assertEqual(document.status, Document.READY, document.error)

    def test_deleting_a_batch_deletes_its_articles_and_scans(self):
        batch = self.batch()
        self.run_batch(batch)
        storage, name = batch.pages.first().file.storage, batch.pages.first().file.name
        self.client.force_login(self.owner)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post(reverse("document_batch", args=[batch.pk]), {"action": "delete"})
        self.assertFalse(Document.objects.exists())
        self.assertFalse(storage.exists(name))


class TidyTests(LibraryTestCase):
    def test_a_new_model_redoes_what_a_model_filled_in_but_not_what_a_person_typed(self):
        document = self.document()
        self.assertEqual(document.auto_fields, ["title", "author", "year", "topics"])
        Document.objects.filter(pk=document.pk).update(author="Typed By Hand", auto_fields=["title", "year", "topics"])
        self.provider.details = {"topics": ["ponds"], "title": "Better title", "author": "Model's guess", "year": 1980}
        with override_settings(DOCUMENT_MODEL="better-model"):
            self.assertIn(document, index.stale_documents())
            index.index_document(document.pk)
            document.refresh_from_db()
            self.assertNotIn(document, index.stale_documents())
        self.assertEqual((document.title, document.author, document.year), ("Better title", "Typed By Hand", 1980))
        self.assertEqual(document.described_with, "better-model")

    def test_an_edit_takes_the_field_back_from_the_model(self):
        document = self.document()
        self.client.force_login(self.owner)
        data = {"title": "Mine", "author": "R. T.", "year": "1979", "visibility": "club", "club": self.club.pk}
        data.update({"topics": ["breeding"], "text": document.text})
        self.client.post(reverse("document_edit", args=[document.pk]), data)
        document.refresh_from_db()
        self.assertNotIn("title", document.auto_fields)
        self.assertIn("author", document.auto_fields)

    def test_a_new_model_rereads_scans_but_not_corrected_ones(self):
        scan = self.document(text=picture(), name="scan.png")
        corrected = self.document(text=picture(400, 301), name="other.png")
        Document.objects.filter(pk=corrected.pk).update(text_edited=True)
        self.assertEqual(scan.read_with, "fake-vision")
        with override_settings(DOCUMENT_MODEL="better-model"):
            stale = set(index.stale_documents())
        self.assertIn(scan, stale)
        self.assertTrue(stale <= {scan, corrected})
        with override_settings(DOCUMENT_MODEL="better-model"):
            self.assertTrue(index._needs_reading(Document.objects.get(pk=scan.pk)))
            self.assertFalse(index._needs_reading(Document.objects.get(pk=corrected.pk)))

    def test_a_blank_the_model_could_not_fill_is_not_asked_about_again(self):
        self.provider.details = {"topics": [], "title": "", "author": "", "year": None}
        document = self.document()
        self.assertNotIn(document, index.stale_documents())

    def test_tidy_is_bounded_and_skips_queued_documents(self):
        for number in range(3):
            self.document(index_it=False, name=f"{number}.md")
        with mock.patch.object(index, "TIDY_PER_NIGHT", 2), mock.patch("auctions.tasks.index_document.delay") as delay:
            self.assertEqual(index.tidy(), 2)
            self.assertEqual(index.tidy(), 1, "the two already queued are skipped")
        self.assertEqual(delay.call_count, 3)


@override_settings(SINGLE_CLUB_MODE=False, ENABLE_CLUB_FINDER=True)
class AccessTests(LibraryTestCase):
    """The library is on per account (``UserData.library_enabled``). Without it, it isn't there at all."""

    TOOLS = {
        "search_documents",
        "read_document",
        "update_document",
        "list_documents",
        "add_document",
        "delete_document",
    }

    def setUp(self):
        super().setUp()
        for user in (self.member, self.staff):
            user.userdata.library_enabled = False
            user.userdata.save(update_fields=["library_enabled"])

    def run_as(self, user, name, **params):
        request = RequestFactory().post("/mcp/")
        request.user = user
        request.palette_page = {}
        return palette_actions.run_action(request, name, params)

    def test_the_pages_are_not_found(self):
        document = self.document(visibility=Visibility.PUBLIC)
        self.client.force_login(self.member)
        for url in (reverse("library"), document.get_absolute_url(), reverse("document_file", args=[document.pk])):
            self.assertEqual(self.client.get(url).status_code, 404, url)
        upload = SimpleUploadedFile("care.txt", b"Feed daphnia twice a day.")
        self.assertEqual(self.client.post(reverse("library_add"), {"file": upload}).status_code, 404)
        self.assertFalse(Document.objects.filter(owner=self.member).exists())

    def test_nothing_links_to_it(self):
        """The menu, the club's sidebar and the clubs guide."""
        sidebar_link = reverse("library") + "?club=" + self.club.slug
        guide = reverse("help_guide", kwargs={"slug": "clubs"})
        club_page = reverse("club_detail", kwargs={"slug": self.club.slug})
        self.client.force_login(self.owner)
        self.assertContains(self.client.get(club_page), sidebar_link)
        self.assertContains(self.client.get(guide), 'id="library"')
        self.owner.userdata.library_enabled = False
        self.owner.userdata.save(update_fields=["library_enabled"])
        for page in (club_page, guide):
            self.assertNotContains(self.client.get(page), reverse("library"), msg_prefix=page)

    def test_the_tools_are_neither_offered_nor_run(self):
        document = self.document(visibility=Visibility.PUBLIC)
        for user in (self.member, self.staff):
            self.assertFalse(self.TOOLS & {action.name for action in palette_actions.actions_for(user)})
        self.assertTrue(self.TOOLS <= {action.name for action in palette_actions.actions_for(self.owner)})
        self.assertIn("error", self.run_as(self.member, "read_document", document=str(document.pk)))
        self.assertIn("error", self.run_as(self.member, "add_document", title="T", text="words"))
        self.assertFalse(Document.objects.filter(owner=self.member).exists())
        self.assertTrue(self.run_as(self.owner, "search_documents", query="microworms")["found"])

    def test_nor_its_recipes_or_its_page_in_the_palette(self):
        self.assertTrue(prompts.LIBRARY_PROMPTS <= set(prompts.BY_NAME))
        offered = {descriptor["name"] for descriptor in prompts.descriptors(self.owner)}
        self.assertTrue(prompts.LIBRARY_PROMPTS <= offered)
        self.assertFalse(
            prompts.LIBRARY_PROMPTS & {descriptor["name"] for descriptor in prompts.descriptors(self.member)}
        )
        self.assertNotIn("library", [route.key for route in palette_routes._permitted_routes(self.member)])
        request = RequestFactory().get("/")
        request.palette_page = {}
        request.user = self.member
        self.assertIn("error", palette_routes.resolve_route(request, palette_routes.ROUTES["library"], {}))
        request.user = self.owner
        self.assertEqual(
            palette_routes.resolve_route(request, palette_routes.ROUTES["library"], {})["url"], reverse("library")
        )

    def test_new_accounts_follow_the_setting_and_a_command_sets_everyone(self):
        for setting in (False, True):
            with override_settings(LIBRARY_ENABLED_FOR_USERS=setting):
                user = User.objects.create_user(f"new{setting}", f"new{setting}@example.com", "x")
            self.assertEqual(user.userdata.library_enabled, setting)
        call_command("change_library", "on", stdout=io.StringIO())
        self.assertFalse(UserData.objects.filter(library_enabled=False).exists())
        call_command("change_library", "off", stdout=io.StringIO())
        self.assertFalse(UserData.objects.filter(library_enabled=True).exists())


def _picture_bytes(kind="JPEG", size=(300, 200), exif=None) -> bytes:
    out = io.BytesIO()
    extra = {"exif": exif} if exif is not None else {}
    Image.new("RGB", size, (20, 120, 200)).save(out, kind, **extra)
    return out.getvalue()


class PictureTests(LibraryTestCase):
    """Pictures on a document: posted to a signed link by an assistant's script or a person's phone."""

    def run_as(self, user, name, **params):
        request = RequestFactory().post("/mcp/")
        request.user = user
        request.palette_page = {}
        return palette_actions.run_action(request, name, params)

    def link(self, document, user=None):
        return pictures.upload_link(document, user or self.owner)

    def setUp(self):
        super().setUp()
        # What curl is: no session and no CSRF token. The link is the credential.
        self.client = Client(enforce_csrf_checks=True)

    def post(self, url, *files, captions=(), **extra):
        data = {"file": [SimpleUploadedFile(name, data) for name, data in files], "caption": list(captions), **extra}
        return self.client.post(url, data)

    def test_add_document_hands_back_a_link_a_script_can_post_to(self):
        result = self.run_as(self.owner, "add_document", title="Killifish eggs", text="Stored in peat.")
        self.assertTrue(result.get("ok"), result)
        self.assertIn("curl", result["how_to_add_pictures"])
        response = self.post(
            result["pictures_upload_url"],
            ("male.jpg", _picture_bytes()),
            ("eggs.png", _picture_bytes("PNG")),
            captions=["Male in colour", "Eggs in peat"],
        )
        self.assertEqual(response.status_code, 200, response.content)
        body = response.json()
        self.assertEqual([each["caption"] for each in body["added"]], ["Male in colour", "Eggs in peat"])
        document = Document.objects.get()
        # Converted as lot images are, by resize_source.
        self.assertEqual([p.file.name[-4:] for p in document.pictures.all()], [".jpg", ".jpg"])

    def test_the_cameras_metadata_is_gone(self):
        exif = Image.Exif()
        exif[0x010F] = "PhoneMaker"  # Make; a GPS block goes the same way
        document = self.document(index_it=False)
        self.post(self.link(document), ("p.jpg", _picture_bytes(exif=exif.tobytes())))
        with document.pictures.get().file.open("rb") as handle, Image.open(handle) as stored:
            self.assertEqual(dict(stored.getexif()), {})

    def test_a_link_is_for_one_document_and_one_keeper(self):
        document = self.document(index_it=False)
        other = self.document(text="other words", name="other.md", index_it=False)
        borrowed = self.link(document).replace(f"/library/{document.pk}/", f"/library/{other.pk}/")
        self.assertEqual(self.post(borrowed, ("p.jpg", _picture_bytes())).status_code, 404)
        self.assertEqual(
            self.post(f"/library/{document.pk}/pictures/forged/", ("p.jpg", _picture_bytes())).status_code, 404
        )
        # Signed by a club admin who has since stopped being one.
        admin = User.objects.create_user("admin2", "admin2@example.com", "x")
        admin.userdata.library_enabled = True
        admin.userdata.save(update_fields=["library_enabled"])
        membership = ClubMember.objects.create(club=self.club, user=admin, name="A", permission_edit_club=True)
        url = self.link(document, admin)
        self.assertEqual(self.client.get(url).status_code, 200)
        membership.permission_edit_club = False
        membership.save()
        self.assertEqual(self.client.get(url).status_code, 404)
        self.assertFalse(document.pictures.exists())

    def test_what_isnt_a_picture_is_refused_and_repeats_are_skipped(self):
        document = self.document(index_it=False)
        response = self.post(self.link(document), ("notes.txt", b"not a picture"))
        self.assertEqual(response.status_code, 400)
        self.assertIn("isn't a picture", response.json()["problems"][0])
        self.assertEqual(self.client.post(self.link(document)).status_code, 400, "no file at all")
        picture = _picture_bytes()
        self.post(self.link(document), ("a.jpg", picture))
        again = self.post(self.link(document), ("a.jpg", picture)).json()
        self.assertIn("already on this document", again["problems"][0])
        with mock.patch.object(pictures, "MAX_PICTURES_PER_DOCUMENT", 1):
            full = self.post(self.link(document), ("b.png", _picture_bytes("PNG"))).json()
        self.assertIn("up to 1 pictures", full["problems"][0])
        self.assertEqual(document.pictures.count(), 1)

    def test_readers_see_pictures_and_only_keepers_get_the_link(self):
        document = self.document(index_it=False)
        self.post(self.link(document), ("a.jpg", _picture_bytes()), captions=["Pair"])
        picture = document.pictures.get()
        theirs = self.run_as(self.member, "read_document", document=str(document.pk))
        self.assertEqual(theirs["pictures"][0]["picture_url"], reverse("document_picture", args=[picture.pk]))
        self.assertNotIn("pictures_upload_url", theirs)
        self.assertIn("pictures_upload_url", self.run_as(self.owner, "read_document", document=str(document.pk)))

        self.client.force_login(self.member)
        self.assertEqual(self.client.get(reverse("document_picture", args=[picture.pk])).status_code, 200)
        page = self.client.get(document.get_absolute_url()).content.decode()
        self.assertIn(reverse("document_picture", args=[picture.pk]), page)
        self.assertNotIn("/pictures/", page, "a reader gets no link to add pictures")
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(reverse("document_picture", args=[picture.pk])).status_code, 404)

    def test_removing_a_picture_and_deleting_the_document_delete_their_files(self):
        document = self.document(index_it=False)
        self.post(self.link(document), ("a.jpg", _picture_bytes()), ("b.png", _picture_bytes("PNG")))
        first, second = document.pictures.all()
        with self.captureOnCommitCallbacks(execute=True):
            link = self.link(document)  # signed with a timestamp: a second call can differ
            response = self.client.post(link, {"remove": first.pk, "from_page": "1"})
        self.assertRedirects(response, link, fetch_redirect_response=False)
        self.assertFalse(first.file.storage.exists(first.file.name))
        with self.captureOnCommitCallbacks(execute=True):
            document.delete()
        self.assertFalse(second.file.storage.exists(second.file.name))

    def test_pictures_count_against_library_space(self):
        document = self.document(index_it=False)
        with mock.patch("auctions.documents.forms.MAX_BYTES_PER_PERSON", document.size + 10):
            body = self.post(self.link(document), ("a.jpg", _picture_bytes())).json()
        self.assertIn("library space", body["problems"][0])


def _picture_on(document, size=(300, 200)):
    added, problems = pictures.add_pictures(
        document, [SimpleUploadedFile("fish.jpg", _picture_bytes(size=size))], ["A male"]
    )
    assert added, problems
    return added[0]


class AddressTests(LibraryTestCase):
    """A document's address carries its title; the old one by number still works, for its readers only."""

    def test_the_title_is_in_the_address_and_a_stale_one_redirects(self):
        document = self.document()
        self.assertEqual(document.get_absolute_url(), f"/library/{document.pk}-dwarf-cichlids/")
        self.client.force_login(self.member)
        self.assertEqual(self.client.get(document.get_absolute_url()).status_code, 200)
        stale = reverse("document_detail", kwargs={"pk": document.pk, "slug": "old-title"})
        self.assertRedirects(self.client.get(stale), document.get_absolute_url())
        self.assertRedirects(self.client.get(f"/library/{document.pk}/"), document.get_absolute_url())

    def test_the_short_link_never_spells_out_a_title_to_somebody_who_cannot_read_it(self):
        document = self.document()
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(f"/library/{document.pk}/").status_code, 404)
        self.client.logout()
        response = self.client.get(f"/library/{document.pk}/")
        self.assertEqual(response.status_code, 302)
        self.assertNotIn("dwarf", response["Location"])


class ListTests(LibraryTestCase):
    """``/library/``: your own and your clubs' documents, filtered by words and by who can read them."""

    def test_it_lists_mine_and_my_clubs_but_not_everyone_elses(self):
        made = []
        for title, kwargs in (
            ("Club notes", {}),
            ("Personal notes", {"club": None, "owner": self.member}),
            ("Elsewhere", {"club": None, "owner": self.outsider, "visibility": Visibility.PUBLIC}),
        ):
            self.provider.details = {"topics": [], "title": title, "author": "", "year": None}
            made.append(self.document(name=f"{title}.md", text=f"# {title}\n\nText.", **kwargs))
        club_one, mine, elsewhere = made
        self.client.force_login(self.member)
        response = self.client.get(reverse("library"))
        self.assertContains(response, club_one.display_title)
        self.assertContains(response, mine.display_title)
        self.assertNotContains(response, elsewhere.display_title)
        self.assertContains(response, reverse("library_add"))
        self.assertContains(response, reverse("library_search"))

    def test_words_and_visibility_filter_over_htmx(self):
        self.provider.details = {"topics": [], "title": "Killifish", "author": "Ann Example", "year": 1977}
        killies = self.document(name="k.md", text="# Killies\n\nPeat.")
        self.provider.details = {"topics": [], "title": "Plants", "author": "Bob", "year": 1990}
        plants = self.document(name="p.md", text="# Plants\n\nJava fern.", club=None, visibility=Visibility.WEB)
        self.client.force_login(self.owner)
        htmx = {"HTTP_HX_REQUEST": "true"}
        response = self.client.get(reverse("library"), {"query": "example 1977"}, **htmx)
        self.assertContains(response, killies.display_title)
        self.assertNotContains(response, plants.display_title)
        self.assertNotContains(response, "<html")
        response = self.client.get(reverse("library"), {"visibility": Visibility.WEB}, **htmx)
        self.assertContains(response, plants.display_title)
        self.assertNotContains(response, killies.display_title)

    def test_old_search_links_land_on_the_search_page_which_still_filters_by_species(self):
        document = self.document()
        document.species.add(self.species)
        self.client.force_login(self.member)
        response = self.client.get(reverse("library"), {"species": self.species.pk, "q": "fry"})
        self.assertRedirects(
            response, reverse("library_search") + f"?species={self.species.pk}&q=fry", fetch_redirect_response=False
        )
        response = self.client.get(reverse("library_search"), {"species": self.species.pk, "q": "fry"})
        self.assertContains(response, document.display_title)
        self.assertContains(response, "Apistogramma agassizii")

    def test_its_own_menu_entry_and_app_drawer_row(self):
        from auctions.mobile.menu import menu_for

        def drawer(user):
            return [row["title"] for section in menu_for(user)["sections"] for row in section["items"]]

        self.client.force_login(self.member)
        self.assertContains(
            self.client.get(reverse("help")), f'<a class="nav-link" href="{reverse("library")}">Library</a>'
        )
        self.assertIn("Library", drawer(self.member))
        self.member.userdata.library_enabled = False
        self.member.userdata.save(update_fields=["library_enabled"])
        self.assertNotContains(self.client.get(reverse("help")), f'href="{reverse("library")}"')
        self.assertNotIn("Library", drawer(self.member))


class AddPageTests(LibraryTestCase):
    """``/library/add/``: files or pasted text, starting from the visibility and club used last."""

    def test_it_starts_public_then_remembers_the_last_choice(self):
        self.client.force_login(self.owner)
        self.assertEqual(
            self.client.get(reverse("library_add")).context["form"].initial,
            {"visibility": Visibility.PUBLIC, "club": None},
        )
        with mock.patch("auctions.tasks.index_document.delay"):
            self.client.post(
                reverse("library_add"),
                {"file": SimpleUploadedFile("a.txt", b"a"), "visibility": Visibility.CLUB, "club": self.club.pk},
            )
        self.assertEqual(
            self.client.get(reverse("library_add")).context["form"].initial,
            {"visibility": Visibility.CLUB, "club": self.club.pk},
        )

    def test_search_engines_are_never_the_remembered_choice(self):
        Document.objects.create(owner=self.owner, visibility=Visibility.WEB, original_name="x")
        self.client.force_login(self.owner)
        self.assertEqual(
            self.client.get(reverse("library_add")).context["form"].initial["visibility"], Visibility.PUBLIC
        )

    def test_a_club_you_can_no_longer_file_under_is_not_remembered(self):
        Document.objects.create(owner=self.member, club=self.club, visibility=Visibility.CLUB, original_name="x")
        self.client.force_login(self.member)
        self.assertEqual(
            self.client.get(reverse("library_add")).context["form"].initial,
            {"visibility": Visibility.PUBLIC, "club": None},
        )

    def test_pasted_text_is_a_document_as_it_was_written(self):
        self.client.force_login(self.owner)
        with mock.patch("auctions.tasks.index_document.delay"):
            response = self.client.post(
                reverse("library_add"), {"text": "# Feeding\r\n\r\nDaphnia, **twice** a day.", "mode": "pages"}
            )
        document = Document.objects.get()
        self.assertRedirects(response, document.get_absolute_url(), fetch_redirect_response=False)
        self.assertEqual(document.text, "# Feeding\n\nDaphnia, **twice** a day.")
        self.assertTrue(document.parser_version)
        self.assertEqual(document.original_name, "pasted text.md")

    def test_nothing_chosen_says_so(self):
        self.client.force_login(self.owner)
        self.assertContains(self.client.post(reverse("library_add"), {}), "Choose a file, or paste the text.")


class WebVisibilityTests(LibraryTestCase):
    """``web``: readable signed out, listed in the sitemap, and the only kind not marked noindex."""

    def test_anyone_reads_a_web_document_and_its_pictures(self):
        document = self.document(visibility=Visibility.WEB)
        picture = _picture_on(document)
        response = self.client.get(document.get_absolute_url())
        self.assertContains(response, "fifth day")
        self.assertContains(response, 'rel="canonical"')
        self.assertNotContains(response, 'name="robots" content="noindex"')
        self.assertNotContains(response, reverse("document_edit", args=[document.pk]))
        self.assertEqual(self.client.get(reverse("document_picture", args=[picture.pk])).status_code, 200)
        self.assertIn(document.get_absolute_url(), self.client.get("/sitemap.xml").content.decode())

    def test_other_documents_still_send_the_signed_out_to_sign_in(self):
        document = self.document(visibility=Visibility.PUBLIC)
        picture = _picture_on(document)
        for url in (document.get_absolute_url(), reverse("document_picture", args=[picture.pk])):
            self.assertRedirects(
                self.client.get(url), f"{reverse('account_login')}?next={url}", fetch_redirect_response=False
            )
        self.assertNotIn(document.get_absolute_url(), self.client.get("/sitemap.xml").content.decode())
        self.client.force_login(self.member)
        self.assertContains(self.client.get(document.get_absolute_url()), 'name="robots" content="noindex"')

    def test_a_scan_shared_with_an_article_not_on_the_web_stays_private(self):
        from auctions.models import BatchPage, DocumentBatch

        batch = DocumentBatch.objects.create(owner=self.owner, visibility=Visibility.CLUB, club=self.club)
        shared = BatchPage.objects.create(batch=batch, file=ContentFile(b"x", name="p.png"), original_name="p.png")
        alone = BatchPage.objects.create(batch=batch, file=ContentFile(b"x", name="q.png"), original_name="q.png")
        web = self.document(name="w.md", visibility=Visibility.WEB, index_it=False)
        club_only = self.document(name="c.md", text="# Other\n\nText.", index_it=False)
        web.pages.add(shared, alone)
        club_only.pages.add(shared)
        self.assertRedirects(
            self.client.get(reverse("batch_page", args=[shared.pk])),
            f"{reverse('account_login')}?next={reverse('batch_page', args=[shared.pk])}",
            fetch_redirect_response=False,
        )
        self.assertEqual(self.client.get(reverse("batch_page", args=[alone.pk])).status_code, 200)

    def test_where_it_came_from_is_not_shown_to_the_open_web(self):
        document = self.document(name="IMG_private_name.md", visibility=Visibility.WEB)
        self.assertNotContains(self.client.get(document.get_absolute_url()), "IMG_private_name")

    def test_an_account_without_the_library_reads_only_web_documents(self):
        web = self.document(name="w.md", visibility=Visibility.WEB)
        public = self.document(name="p.md", text="# Other\n\nText.", visibility=Visibility.PUBLIC)
        self.member.userdata.library_enabled = False
        self.member.userdata.save(update_fields=["library_enabled"])
        self.client.force_login(self.member)
        self.assertEqual(self.client.get(web.get_absolute_url()).status_code, 200)
        self.assertEqual(self.client.get(public.get_absolute_url()).status_code, 404)


class MarkdownTests(LibraryTestCase):
    """The text is Markdown, cut down to safe HTML, with the document's own pictures placed in it."""

    def test_markdown_renders_and_scripts_and_other_sites_pictures_do_not(self):
        from auctions.documents.render import to_html

        html = to_html(
            "## Fry\n\n**Daphnia** <script>alert(1)</script> [link](javascript:alert(1)) "
            "![tracker](https://example.com/x.png) <img src=x onerror=alert(1)>",
            [],
        )
        self.assertIn("<h2>Fry</h2>", html)
        self.assertIn("<strong>Daphnia</strong>", html)
        self.assertNotIn("<script", html)
        self.assertNotIn("javascript:", html)
        self.assertNotIn("example.com", html)
        self.assertNotIn("<img", html)
        self.assertIn("tracker", html)
        # Angle brackets in an old transcription are words, not tags to throw away.
        self.assertIn("&lt;Betta splendens&gt;", to_html("Keep <Betta splendens> warm", []))

    def test_a_passage_cut_out_of_a_long_table_is_rendered_with_the_rest_of_it(self):
        rows = "\n".join(f"| {n} | microworms and brine shrimp, fed twice a day |" for n in range(120))
        document = self.document(text=f"# Feeding log\n\n| Day | Food |\n|---|---|\n{rows}\n")
        self.assertGreater(document.chunks.count(), 2)
        self.client.force_login(self.member)
        response = self.client.get(document.get_absolute_url())
        self.assertNotContains(response, "| 119 |")
        self.assertContains(response, "<td>119</td>")
        for start in document.chunks.values_list("start", flat=True):
            self.assertContains(response, f'id="at-{start}"')

    def test_a_picture_the_text_names_is_shown_there_and_not_again_above(self):
        document = self.document(visibility=Visibility.PUBLIC)
        placed = _picture_on(document)
        other_document = self.document(name="o.md", text="# Other\n\nText.")
        foreign = _picture_on(other_document)
        url = reverse("document_picture", args=[placed.pk])
        foreign_url = reverse("document_picture", args=[foreign.pk])
        document.text += f"\n\n![The male]({url})\n\n![Not ours]({foreign_url})\n"
        document.chunker_version = ""
        document.save()
        index.index_document(document.pk)
        document.refresh_from_db()
        self.client.force_login(self.member)
        response = self.client.get(document.get_absolute_url())
        self.assertEqual(response.context["gallery"], [])
        self.assertContains(response, f'src="{url}"', count=1)
        self.assertNotContains(response, f'src="{foreign_url}"')
        self.assertContains(response, 'id="picture-zoom"')
        self.assertNotContains(response, 'target="_blank" rel="noopener"><img')

    def test_pictures_are_capped_at_a_thousand_pixels(self):
        document = self.document()
        picture = _picture_on(document, size=(2400, 1200))
        with picture.file.open("rb") as handle, Image.open(handle) as image:
            self.assertEqual(image.size, (1000, 500))


@override_settings(DOCUMENT_SERVER_OCR=False)
class TranscribeTests(LibraryTestCase):
    """The site reads no pictures: scans wait for a person's agent, which works through them over ``/mcp/``."""

    def run_as(self, user, name, **params):
        request = RequestFactory().post("/mcp/")
        request.user = user
        request.palette_page = {}
        return palette_actions.run_action(request, name, params)

    def scan(self, **kwargs) -> Document:
        with mock.patch("auctions.tasks.index_document.delay"):
            return self.document(text=picture(1200, 1600), name="page.png", **kwargs)

    def scanned_batch(self, pages=3):
        from auctions.documents import stitch
        from auctions.models import BatchPage, DocumentBatch

        batch = DocumentBatch.objects.create(owner=self.owner, club=self.club, visibility=Visibility.CLUB)
        for number in range(pages, 0, -1):
            page = BatchPage(batch=batch, original_name=f"scan{number}.jpg")
            page.file.save(f"scan{number}.jpg", ContentFile(picture(800, 1000, fmt="JPEG")), save=True)
        stitch.process(batch.pk)
        batch.refresh_from_db()
        return batch

    def test_a_scan_waits_and_the_site_pays_for_nothing(self):
        document = self.scan()
        self.assertEqual(document.status, Document.TRANSCRIBE)
        self.assertEqual(document.text, "")
        self.assertEqual(self.provider.calls["complete"], 0)
        self.assertNotIn(document, index.stale_documents(), "the nightly tidy leaves it alone")

    def test_text_files_are_still_read(self):
        self.assertEqual(self.document().status, Document.READY)

    def test_a_reading_from_before_stands(self):
        document = self.scan()
        Document.objects.filter(pk=document.pk).update(text="Read by the old model.", parser_version="old")
        index.index_document(document.pk)
        document.refresh_from_db()
        self.assertEqual((document.status, document.text), (Document.READY, "Read by the old model."))
        self.assertEqual(self.provider.calls["complete"], 0)

    def test_only_who_may_edit_it_sees_it_waiting(self):
        document = self.scan(visibility=Visibility.PUBLIC)
        reference = f"document {document.pk}"
        for user, sees in ((self.owner, True), (self.staff, True), (self.member, False), (self.outsider, False)):
            listed = [row["item"] for row in self.run_as(user, "waiting_articles")["articles"]]
            self.assertEqual(reference in listed, sees, user)
        self.assertIn("error", self.run_as(self.member, "read_waiting_page", item=reference))
        self.assertIn("error", self.run_as(self.member, "save_transcription", item=reference, text="x"))

    def test_an_agent_reads_a_page_and_saves_it_for_the_uploader(self):
        document = self.scan(visibility=Visibility.PUBLIC)
        reference = f"document {document.pk}"
        page = self.run_as(self.staff, "read_waiting_page", item=reference)
        self.assertEqual(page["pages"], 1)
        (jpeg,) = page[palette_actions.KEY_IMAGES]
        self.assertEqual(jpeg[:2], b"\xff\xd8")
        self.assertIn("error", self.run_as(self.staff, "read_waiting_page", item=reference, page=2))
        with mock.patch("auctions.tasks.index_document.delay", side_effect=index.index_document):
            with self.captureOnCommitCallbacks(execute=True):
                result = self.run_as(
                    self.staff,
                    "save_transcription",
                    item=reference,
                    text="<!-- page 1 -->\n\nThe Apistogramma agassizii spawned in a clay pot.",
                    title="Apistos",
                    year=1978,
                )
        self.assertTrue(result.get("ok"), result)
        document.refresh_from_db()
        self.assertEqual(
            (document.owner, document.club, document.visibility), (self.owner, self.club, Visibility.PUBLIC)
        )
        self.assertEqual((document.title, document.year, document.status), ("Apistos", 1978, Document.READY))
        self.assertTrue(document.text_edited, "never read over again")
        self.assertIn(self.species, document.species.all())
        self.assertEqual(self.run_as(self.owner, "waiting_articles")["articles"], [])

    def test_a_batch_of_scans_waits_and_is_saved_article_by_article(self):
        batch = self.scanned_batch()
        self.assertEqual(batch.status, batch.TRANSCRIBE)
        self.assertEqual(self.provider.calls, {"complete": 0, "complete_json": 0, "embed": 0})
        (row,) = self.run_as(self.owner, "waiting_articles")["articles"]
        self.assertEqual((row["item"], row["pages"], row["pages_already_in_an_article"]), (f"batch {batch.pk}", 3, []))
        names = list(batch.pages.order_by("position").values_list("original_name", flat=True))
        self.assertEqual(names, ["scan1.jpg", "scan2.jpg", "scan3.jpg"], "numbered in file-name order")
        item = f"batch {batch.pk}"
        self.assertIn("more_info_needed", self.run_as(self.owner, "save_transcription", item=item, text="x", title="T"))
        self.assertIn(
            "more_info_needed",
            self.run_as(self.owner, "save_transcription", item=item, text="x", title="T", pages=[4]),
        )
        first = self.run_as(self.owner, "save_transcription", item=item, text="Bettas.", title="Bettas", pages=[1, 2])
        self.assertTrue(first.get("ok"), first)
        again = self.run_as(self.owner, "save_transcription", item=item, text="Bettas!", title="Bettas", pages=[1, 2])
        self.assertTrue(again.get("ok"), again)
        self.assertEqual(batch.documents.get().text, "Bettas!", "saved again, not twice")
        (row,) = self.run_as(self.owner, "waiting_articles")["articles"]
        self.assertEqual(row["pages_already_in_an_article"], [1, 2])
        last = self.run_as(
            self.owner, "save_transcription", item=item, text="Killifish.", title="Killies", pages="3", finished=True
        )
        self.assertTrue(last.get("ok"), last)
        batch.refresh_from_db()
        self.assertEqual(batch.status, batch.DONE)
        killies = batch.documents.get(title="Killies")
        self.assertEqual((killies.owner, killies.club, killies.visibility), (self.owner, self.club, Visibility.CLUB))
        self.assertEqual(self.run_as(self.owner, "waiting_articles")["articles"], [])

    def test_the_page_reaches_the_agent_as_a_picture_on_either_endpoint(self):
        from auctions.mcp import admin

        document = self.scan()
        request = RequestFactory().post("/mcp/")
        request.user = self.owner
        result = tools.call_tool(request, "read_waiting_page", {"item": f"document {document.pk}"})
        images = [block for block in result["content"] if block["type"] == "image"]
        self.assertEqual(len(images), 1)
        self.assertEqual(images[0]["mimeType"], "image/jpeg")
        self.assertNotIn("_images", result["content"][0]["text"])
        kept = admin.private_result(result, "read_waiting_page")["content"]
        self.assertEqual([block["type"] for block in kept], ["text", "image"])
        dropped = admin.private_result(result, "describe_lot")["content"]
        self.assertEqual([block["type"] for block in dropped], ["text"], "no other tool's pictures leave")

    def test_the_pages_say_it_waits_and_offer_an_agent(self):
        document = self.scan()
        self.client.force_login(self.owner)
        self.assertContains(self.client.get(document.get_absolute_url()), reverse("user_api_keys"))
        self.assertContains(self.client.get(reverse("library")), "1 article waiting to be transcribed")
        self.client.force_login(self.member)
        response = self.client.get(document.get_absolute_url())
        self.assertContains(response, "Waiting to be transcribed")
        self.assertNotContains(response, reverse("user_api_keys"))

    def test_nobody_else_sees_an_only_me_scan(self):
        document = self.scan(visibility=Visibility.PRIVATE)
        self.assertEqual(len(self.run_as(self.owner, "waiting_articles")["articles"]), 1)
        self.assertEqual(self.run_as(self.staff, "waiting_articles")["articles"], [], "as on the web: a 404")
        self.assertIn("error", self.run_as(self.staff, "read_waiting_page", item=f"document {document.pk}"))

    def test_the_text_under_another_name_is_still_the_text(self):
        batch = self.scanned_batch(pages=1)
        result = self.run_as(
            self.owner, "save_transcription", batch=str(batch.pk), content="Bettas.", title="B", pages=[1]
        )
        self.assertIn("error", result, "a bare number isn't an item")
        result = self.run_as(
            self.owner,
            "save_transcription",
            batch=f"batch {batch.pk}",
            content="Bettas.",
            title="B",
            pages=[1],
            finished=True,
        )
        self.assertTrue(result.get("ok"), result)
        self.assertEqual(batch.documents.get().text, "Bettas.", "never closed with the article thrown away")

    def test_pages_added_later_keep_the_numbers_already_given(self):
        from auctions.documents import stitch
        from auctions.models import BatchPage

        batch = self.scanned_batch(pages=2)
        page = BatchPage(batch=batch, original_name="scan1b.jpg")
        page.file.save("scan1b.jpg", ContentFile(picture(800, 1000, fmt="JPEG")), save=True)
        batch.status = batch.READING
        batch.save(update_fields=["status"])
        stitch.process(batch.pk)
        batch.refresh_from_db()
        self.assertEqual(batch.status, batch.TRANSCRIBE)
        names = list(batch.pages.order_by("position").values_list("original_name", flat=True))
        self.assertEqual(names, ["scan1.jpg", "scan2.jpg", "scan1b.jpg"])

    def test_turning_the_sites_reading_on_picks_up_what_waited(self):
        from auctions.documents import stitch

        untouched, started = self.scanned_batch(), self.scanned_batch()
        self.run_as(self.owner, "save_transcription", item=f"batch {started.pk}", text="B.", title="B", pages=[1])
        with override_settings(DOCUMENT_SERVER_OCR=True), mock.patch.object(stitch, "_send"):
            stitch.requeue_stuck()
        untouched.refresh_from_db()
        started.refresh_from_db()
        self.assertEqual(untouched.status, untouched.READING)
        self.assertFalse(untouched.pages.filter(read=True).exists())
        self.assertEqual(started.status, started.TRANSCRIBE, "never rebuilt over an agent's articles")

    def test_the_owners_routine_proposes_and_the_approval_saves_it(self):
        from auctions.mcp import admin
        from auctions.models import AgentProposal

        document = self.scan(owner=self.member, club=self.club)  # club-only: staff isn't in the club
        mine = self.scan(owner=self.staff, club=None, visibility=Visibility.PUBLIC)
        item = f"document {document.pk}"
        self.assertEqual(self.run_as(self.staff, "waiting_articles")["articles"][0]["item"], f"document {mine.pk}")
        self.assertIn("error", self.run_as(self.staff, "save_transcription", item=item, text="x"), "not on /mcp/")
        request = RequestFactory().post("/mcp/admin/")
        request.user, request.palette_page, request.owner_reach = self.staff, {}, True
        listed = palette_actions.run_action(request, "waiting_articles", {})["articles"]
        self.assertEqual(
            {(row["item"], row["yours"]) for row in listed}, {(item, False), (f"document {mine.pk}", True)}
        )
        steps = [{"tool": "save_transcription", "arguments": {"item": item, "text": "Read."}}]
        proposed = admin.propose_change(request, {"summary": "Transcribe a scan", "steps": steps})
        self.assertTrue(proposed.get("ok"), proposed)
        again = admin.propose_change(request, {"summary": "Transcribe it again", "steps": steps})
        self.assertIn("already transcribes", again["error"])
        self.assertIn(
            "error",
            admin.propose_change(
                request,
                {
                    "summary": "s",
                    "steps": [{"tool": "save_transcription", "arguments": {"item": "document 999999", "text": "x"}}],
                },
            ),
        )
        self.assertEqual(Document.objects.get(pk=document.pk).text, "", "nothing changes until approved")
        approver = RequestFactory().post("/admin-dashboard/proposals/")
        approver.user = self.staff
        proposal = AgentProposal.objects.get()
        admin.apply_proposal(proposal, approver)
        self.assertEqual(proposal.status, AgentProposal.STATUS_APPLIED, proposal.results)
        document.refresh_from_db()
        self.assertEqual((document.text, document.owner, document.visibility), ("Read.", self.member, Visibility.CLUB))
